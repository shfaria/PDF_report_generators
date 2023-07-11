from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle, TA_CENTER, TA_LEFT
from reportlab.lib.units import inch, mm, cm
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Image
from reportlab.platypus.tables import Table,TableStyle,colors
from reportlab.lib.pagesizes import letter, landscape, A4
from reportlab.lib.enums import TA_JUSTIFY
from django.shortcuts import render
from django.http import HttpResponse
from io import BytesIO
import time
from app.models import demoModel
from io import BytesIO, StringIO
import pdftotext
from openpyxl import Workbook
import pdfplumber

#This is your data collected from your Vizard experiment
subject1 = 'Tom'
subject2 = 'Ana'
results1 = [15,23,42,56,76]
results2 = [34,67,94,31,56]

#take the data and make ready for paragraph
def dataToParagraph(name, data):
   
    p = '<strong>Subject name: </strong>' + name + '<br/>' + '<strong>Data: </strong>  ('
    for i in range(len(data)):
        p += str(data[i])
        if i != len(data) - 1:
            p += ', '
        else:
            p += ')'   
    return p

#take the data and convert to list of strings ready for table
def dataToTable(name, data):
   
    data = [str(x) for x in data]
    data.insert(0,name)
    return data


#create the table for our document
def myTable(tabledata):

    #first define column and row size
    colwidths = (70, 50, 50, 50, 50, 50)
    rowheights = (25, 20, 20)

    t = Table(tabledata, colwidths, rowheights)

    GRID_STYLE = TableStyle(
    [('GRID', (0,0), (-1,-1), 0.25, colors.black),
    ('ALIGN', (1,1), (-1,-1), 'RIGHT')]
    )

    t.setStyle(GRID_STYLE)
    return t

class Test(object):
    def __init__(self):
        """Constructor"""
        self.width, self.height = A4
        self.styles = getSampleStyleSheet()
        

    def coord(self, x, y, unit=1):
        x, y = x * unit, self.height -  y * unit
        return x, y
        

    
    def createDocument(self, canvas, doc):
        """
        Create the document
        """
        self.c = canvas
        normal = self.styles["Normal"]
        centered = ParagraphStyle(name="centered", alignment=TA_CENTER)
        
        header_text = "<b>This is a test header</b>"
        p = Paragraph(header_text, normal)
        p.wrapOn(self.c, self.width, self.height)
        p.drawOn(self.c, *self.coord(80, 12, mm))
        # p.drawOn(self.c, *self.coord(1.8, 9.6, cm))
        # p.drawOn(self.c, )

        ptext = """Lorem ipsum dolor sit amet, consectetur adipisicing elit,
        sed do eiusmod tempor incididunt ut labore et dolore magna aliqua. 
        Ut enim ad minim veniam, quis nostrud exercitation ullamco laboris 
        nisi ut aliquip ex ea commodo consequat. Duis aute irure dolor in
        reprehenderit in voluptate velit esse cillum dolore eu fugiat nulla
        pariatur. Excepteur sint occaecat cupidatat non proident, sunt in 
        culpa qui officia deserunt mollit anim id est laborum."""
        
        p = Paragraph(ptext, style=normal)
        p.wrapOn(self.c, self.width-50, self.height)
        p.drawOn(self.c, 30, 700)
        
        ptext = """
        At vero eos et accusamus et iusto odio dignissimos ducimus qui 
        blanditiis praesentium voluptatum deleniti atque corrupti quos dolores 
        et quas molestias excepturi sint occaecati cupiditate non provident, 
        similique sunt in culpa qui officia deserunt mollitia animi, id est laborum
        et dolorum fuga. Et harum quidem rerum facilis est et expedita distinctio. 
        Nam libero tempore, cum soluta nobis est eligendi optio cumque nihil impedit
        quo minus id quod maxime placeat facere possimus, omnis voluptas assumenda est,
        omnis dolor repellendus. Temporibus autem quibusdam et aut officiis debitis aut 
        rerum necessitatibus saepe eveniet ut et voluptates repudiandae sint et 
        molestiae non recusandae. Itaque earum rerum hic tenetur a sapiente delectus,
        ut aut reiciendis voluptatibus maiores alias consequatur aut perferendis
        doloribus asperiores repellat.
        """
        p = Paragraph(ptext, style=normal)
        p.wrapOn(self.c, self.width-50, self.height)
        p.drawOn(self.c, 30, 600)
        

    def createLineItems(self):
        """
        Create the line items
        """
        text_data = ["Line", "DOS", "Procedure","Description", "Units", "harges","Type1", "Type2","Type3", "Allowance", "demo", "demo2"]
        d = []
        font_size = 8
        centered = ParagraphStyle(name="centered", alignment=TA_CENTER)
        for text in text_data:
            ptext = '<font size="%s"><b>%s</b></font>' % (font_size, text)
            p = Paragraph(ptext, centered)
            d.append(p)
        
        data = [d]
        # data.insert(1, ['']) 
        line_num = 1
        formatted_line_data = [] 

        dem = demoModel.objects.values_list()[250:]
        demo = list(dem)
        for x in demo:
            line_data = list(x)
            line_data.append("12")
            line_data.append("123")

        # for x in range(200):
        #     line_data = [str(line_num), "04/12/13", "73090", "Test Reflexes", "1", "131.00", "0.00", "0.00", "0.00", "0.00", 1233, 4555]
            for item in line_data:
                ptext = '<font size="%s">%s</font>' % (font_size-1, item)
                p = Paragraph(ptext, centered)
                formatted_line_data.append(p)
            data.append(formatted_line_data)
            formatted_line_data = []
            line_num += 1
        
        table = Table(data, repeatRows=1)

        # LIST_STYLE = TableStyle(
        #     [('LINEABOVE', (0,0), (-1,0), 2, colors.green),
        #     ('LINEABOVE', (0,1), (-1,-1), 0.25, colors.black),
        #     ('LINEBELOW', (0,-1), (-1,-1), 2, colors.green),
        #     ('ALIGN', (1,1), (-1,-1), 'RIGHT')]
        #     )
        LIST_STYLE = TableStyle(
            [
            ('LINEBELOW', (0,0), (-1,0), 1, colors.black),
            ('BOTTOMPADDING', (0,0), (-1,0), 3),
            ]
            )
        table.setStyle(LIST_STYLE)
        self.story.append(table)


    def run(self):
        # """
        # Run the report
        # """
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="hello.pdf"'
        # buffer = BytesIO()
        # self.doc = SimpleDocTemplate(response,pagesize=A4, rightMargin=72,leftMargin=72, topMargin=72,bottomMargin=18)
        # # self.doc = SimpleDocTemplate(response)
        # self.story = [Spacer(1, 2.5*inch)]
        # self.createLineItems()
        # self.doc.build(self.story, onFirstPage=self.createDocument)

        # buffer.write(response.content)
        # buffer.seek(0)
        # pdf = pdfplumber.open(buffer, laparams = { "line_overlap": 0.7 })
        # p0 = pdf.pages[0].extract_words()
        # # print(p0)
        # first_page = pdf.pages[0]
        # first_lines = first_page.extract_text(layout=True)
        # # print(first_lines.split("\t"))
        # # for i in first_lines.split("\n"):
        # #     print(i)
        # print(first_lines.split("\n"))
     
        # table = first_page.extract_table(table_settings={
        #     "vertical_strategy": "text",
        #     "horizontal_strategy": "text",
        # })
        # table = [row for row in table if ''.join([str(i) for i in row]) != '']
        # print(table)
        with open("static/hellopy.pdf", "rb") as f:
            pdf = pdftotext.PDF(f, physical=True)
        page1 = pdf[0].split("\n")
        print(page1)
        wb = Workbook(write_only=True)
        ws = wb.create_sheet()
        ws.sheet_view.showGridLines = False
        
        for line in page1:
            if line == '\x0c':
                continue
            line = line.split("\t")
            line = list(line)
            # print(line)
            ws.append(line)
        # for eachline in first_lines
        # print(first_lines.split("\n"))

        wb.save('static/plum.xlsx')

        # buffer.seek(0)

        return response

        

if __name__ == "__main__":
    t = Test()
    t.run()


def demo(request):
    t = Test()
    result = t.run()
    return result

