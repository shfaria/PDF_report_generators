from django.shortcuts import render
from io import BytesIO
import time
from reportlab.pdfgen import canvas
from django.http import HttpResponse
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.pagesizes import A4
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib.units import inch
from reportlab.platypus.tables import Table,TableStyle,colors
import csv
from django.http import StreamingHttpResponse, HttpResponse
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch, mm
from reportlab.platypus import SimpleDocTemplate, Spacer, Paragraph, Table, TableStyle

from reportlab.pdfgen import canvas
from io import BytesIO, StringIO
import pandas as pd
from pypdf import PdfReader


def xlreport(**kwargs):
    columns = kwargs['headers']
    table_values = kwargs['table_values']
    filename = str(kwargs['filename'])

    excel_data = [columns, *table_values]
    if excel_data:
        wb = Workbook(write_only=True)
        ws = wb.create_sheet()
        for line in excel_data:
            if type(line) != list:
                line = list(line)
            ws.append(line)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=%s' % filename
    wb.save(response)
    return response

